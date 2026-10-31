"""TDD: companion resume renders as a 1-page compact document.

These tests inspect the generated .docx (no Word/PDF required) so they can
fail before the compact path exists.
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import make_resume as mr  # noqa: E402


def _para_text(p) -> str:
    from docx.oxml.ns import qn

    chunks = []
    for child in p._p:
        if child.tag == qn("w:r"):
            chunks.append("".join(t.text or "" for t in child.findall(qn("w:t"))))
        elif child.tag == qn("w:hyperlink"):
            chunks.append("".join(t.text or "" for t in child.findall(".//" + qn("w:t"))))
    return "".join(chunks)


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        parts.append(_para_text(p))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    parts.append(_para_text(p))
    for section in doc.sections:
        for p in section.footer.paragraphs:
            parts.append(_para_text(p))
    return "\n".join(parts)


def _hyperlink_targets(path: Path) -> list[str]:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    found = []

    def walk(element, rels):
        hrefs = {
            rel.rId: rel.target_ref
            for rel in rels.values()
            if "hyperlink" in rel.reltype
        }
        for hl in element.findall(".//" + qn("w:hyperlink")):
            rid = hl.get(qn("r:id"))
            if rid and rid in hrefs:
                found.append(hrefs[rid])

    walk(doc.element.body, doc.part.rels)
    for section in doc.sections:
        walk(section.footer._element, section.footer.part.rels)
    return found


def _inline_shape_count(path: Path) -> int:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    n = len(doc.inline_shapes)
    for section in doc.sections:
        n += len(section.footer._element.findall(".//" + qn("w:drawing")))
    return n


class CompactResumeTests(unittest.TestCase):
    def setUp(self):
        self.timeline, self.meta = mr.load_data()
        self._tmp = tempfile.TemporaryDirectory()
        self.out = Path(self._tmp.name) / "resume.docx"
        self._orig_output = mr.OUTPUT
        mr.OUTPUT = self.out

    def tearDown(self):
        mr.OUTPUT = self._orig_output
        self._tmp.cleanup()

    def test_load_data_exposes_resume_only_fields(self):
        self.assertIn("resumeSummary", self.meta)
        self.assertTrue(self.meta["resumeSummary"].strip())
        self.assertLess(len(self.meta["resumeSummary"]), len(self.meta["summary"]))
        self.assertIn("resumeCompetencies", self.meta)
        self.assertGreaterEqual(len(self.meta["resumeCompetencies"]), 6)
        self.assertLessEqual(len(self.meta["resumeCompetencies"]), len(self.meta["competencies"]))
        gemmacon = next(i for i in self.timeline if i["org"].startswith("GEMMACON"))
        self.assertEqual(len(gemmacon["resumeBullets"]), 3)
        paccar = next(i for i in self.timeline if i["org"] == "PACCAR")
        self.assertEqual(len(paccar["resumeBullets"]), 1)
        edu = next(i for i in self.timeline if i.get("type") == "education")
        self.assertEqual(edu.get("resumeBullets"), [])

    def test_compact_omits_kytimes_and_competency_table(self):
        mr.build_document(self.timeline, self.meta, compact=True)
        text = _docx_text(self.out)
        self.assertNotIn("kytimes", text.lower())
        self.assertNotIn("Knew York", text)
        self.assertNotIn("Agentic AI Systems & Workflows", text)
        self.assertIn("Director of AI Deployment & Solutions", text)
        self.assertIn("MIT xPRO", text)
        self.assertRegex(text, r"(?i)companion site")
        self.assertGreaterEqual(_inline_shape_count(self.out), 2)
        self.assertIn("Mercedes-Benz R&D North America", text)
        self.assertIn(self.meta["resumeSummary"][:40], text)
        self.assertIn(self.meta["name"], text)
        self.assertNotIn("haven't articulated", text)
        self.assertNotRegex(text, r"(?i)technical pre[-\s]?sales")

    def test_compact_web_urls_are_hyperlinks(self):
        mr.build_document(self.timeline, self.meta, compact=True)
        urls = _hyperlink_targets(self.out)
        self.assertIn("mailto:" + self.meta["email"], urls)
        self.assertIn("https://" + self.meta["linkedin"], urls)
        self.assertIn("https://" + self.meta["x"], urls)
        self.assertIn("https://" + self.meta["github"], urls)
        self.assertIn(mr.COMPANION_URL, urls)
        self.assertIn(mr.CERT_URL, urls)
        self.assertTrue(mr.CERT_URL.endswith("#acc.3MJw8mIU"))
        # QR image + caption both point at the companion site.
        self.assertGreaterEqual(urls.count(mr.COMPANION_URL), 2)
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(str(self.out))
        colors, underlines = set(), set()
        for hl in doc.element.body.findall(".//" + qn("w:hyperlink")):
            for color in hl.findall(".//" + qn("w:color")):
                colors.add(color.get(qn("w:val")))
            for u in hl.findall(".//" + qn("w:u")):
                underlines.add(u.get(qn("w:val")))
        self.assertIn("0563C1", colors)
        self.assertIn("single", underlines)

    def test_experience_copy_is_plain_text_without_headcount_claim(self):
        for compact in (True, False):
            mr.build_document(self.timeline, self.meta, compact=compact)
            text = _docx_text(self.out)
            self.assertNotIn("300%", text)
            self.assertNotIn("headcount", text.lower())
            self.assertNotIn("](http", text)
            self.assertNotIn("hyundainews.com", text)
            self.assertNotIn("newsroom.porsche.com", text)
            self.assertNotIn("toyotaconnected.com", text)
            self.assertNotIn("prnewswire.com", text)
            self.assertIn("Promoted twice in two years", text)
            self.assertIn("all-electric Macan", text)
            self.assertIn("from 7 to 35 employees", text)
            self.assertIn("Supported QA for Hyundai Blue Link", text)

    def test_plain_bullet_strips_markdown_links(self):
        self.assertEqual(
            mr.plain_bullet("Named an [Engadget finalist](https://example.com/a), then shipped."),
            "Named an Engadget finalist, then shipped.",
        )
        self.assertNotIn("http", mr.plain_bullet("See https://example.com/path for details."))

    def test_default_build_still_includes_projects_for_job_search(self):
        mr.build_document(self.timeline, self.meta)
        text = _docx_text(self.out)
        self.assertIn("Knew York", text)
        self.assertIn("Agentic AI Systems", text)


if __name__ == "__main__":
    unittest.main()
